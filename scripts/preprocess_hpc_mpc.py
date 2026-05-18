#!/usr/bin/env python3
"""
Parse HPC/MPC cohort classification from Excel into JSON.
Maps Excel reason names to DB return_reason values.
"""
import json
import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.expanduser('~/Downloads/Central Sheet - HPC _ MPC Cohorts.xlsx')
OUT_PATH = os.path.join(BASE_DIR, 'data', 'processed', 'hpc_mpc_data.json')

REASON_MAP = {
    'Damage': 'DAMAGED_PRODUCT',
    'Misshipment - Mismatch': 'MISSHIPMENT',
}

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

lookup = {}

# HPC sheet: (vertical, reason) → HPC
ws = wb['HPC']
for row in ws.iter_rows(min_row=1, values_only=True):
    vert, reason = row[0], row[1]
    if not vert or not reason:
        continue
    db_reason = REASON_MAP.get(reason)
    if db_reason:
        lookup[f"{vert}|{db_reason}"] = 'HPC'

# MPC sheet: (vertical, "Misshipment and Damage") → MPC for both reasons
ws = wb['MPC']
for row in ws.iter_rows(min_row=2, values_only=True):
    vert, reason = row[0], row[1]
    if not vert:
        continue
    for db_reason in ['DAMAGED_PRODUCT', 'MISSHIPMENT']:
        key = f"{vert}|{db_reason}"
        if key not in lookup:
            lookup[key] = 'MPC'

output = {
    'description': 'HPC/MPC cohort classification. Key = "vertical|return_reason"',
    'lookup': lookup,
    'stats': {
        'hpc_entries': sum(1 for v in lookup.values() if v == 'HPC'),
        'mpc_entries': sum(1 for v in lookup.values() if v == 'MPC'),
    },
}

with open(OUT_PATH, 'w') as f:
    json.dump(output, f, indent=2)

print(f"HPC entries: {output['stats']['hpc_entries']}")
print(f"MPC entries: {output['stats']['mpc_entries']}")
print(f"Total: {len(lookup)}")
print(f"Output: {OUT_PATH}")
